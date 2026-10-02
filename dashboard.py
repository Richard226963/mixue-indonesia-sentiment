import io
import json
import os

import matplotlib
matplotlib.use("agg")
import matplotlib.pyplot as plt
import pandas as pd
import requests
import streamlit as st
from wordcloud import WordCloud

st.set_page_config(page_title="Analisis Sentimen Mixue Indonesia", layout="wide")

ORDER = ["positive", "neutral", "negative"]
COLORS = {"positive": "#2e9e5b", "neutral": "#9a9a9a", "negative": "#d64545"}
DATA_FILES = ["results.csv", "clean.csv", "metrics.csv", "aspect_summary.csv",
              "news_articles.csv", "news_comments.csv"]

RUN_STEPS = """python scrape.py      # instagram, login sendiri
python news.py        # suara.com
python clean.py       # preprocessing + filter bot
python sentiment.py   # labeling indoBERT + evaluasi
python absa.py
python topics.py
python analytics.py"""


def read_csv(path):
    if not os.path.exists(path):
        return None
    try:
        return pd.read_csv(path)
    except Exception:
        return None


def sidebar_status():
    st.sidebar.subheader("status data")
    for f in DATA_FILES:
        ada = os.path.exists(f"data/{f}")
        st.sidebar.write(f"{'ada ' if ada else 'belum'} - data/{f}")
    st.sidebar.caption("jalankan pipeline dulu kalau masih ada yang belum")


def trend_frame(df):
    dates = pd.to_datetime(df["date"], errors="coerce", format="mixed", utc=True)
    dates = dates.dt.tz_convert("Asia/Jakarta").dt.tz_localize(None)
    rule = "D"
    for r in ("Y", "M", "W", "D"):
        if dates.dt.to_period(r).dropna().nunique() > 1:
            rule = r
            break
    return df.assign(p=dates.dt.to_period(rule).astype(str)) \
             .groupby(["p", "label"]).size().unstack(fill_value=0) \
             .reindex(columns=ORDER, fill_value=0)


def page_dashboard():
    st.title("Analisis Sentimen Mixue Indonesia")
    st.caption("komentar instagram (akun resmi + hashtag) dan portal berita suara.com, "
               "difilter dari bot lalu dilabeli indoBERT secara lokal")

    df = read_csv("data/results.csv")
    if df is None:
        st.warning("belum ada data/results.csv - jalankan pipeline dulu:")
        st.code(RUN_STEPS, language="bash")
        return

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("komentar organik", f"{len(df):,}")
    c2.metric("positif", f"{(df['label'] == 'positive').mean():.0%}")
    c3.metric("negatif", f"{(df['label'] == 'negative').mean():.0%}")
    clean = read_csv("data/clean.csv")
    bots = int(clean["is_bot"].sum()) if clean is not None and "is_bot" in clean else 0
    c4.metric("bot/spam difilter", f"{bots:,}")

    left, right = st.columns(2)
    with left:
        st.subheader("distribusi sentimen")
        counts = df["label"].value_counts().reindex(ORDER).fillna(0)
        st.bar_chart(counts)
    with right:
        st.subheader("instagram vs portal berita")
        share = df.groupby(["source", "label"]).size().unstack(fill_value=0) \
                  .reindex(columns=ORDER, fill_value=0)
        share = share.div(share.sum(axis=1), axis=0)
        st.bar_chart(share)

    st.subheader("tren sentimen")
    counts_t = trend_frame(df)
    if len(counts_t) > 1:
        fig, ax = plt.subplots(figsize=(9.5, 3.4), constrained_layout=True)
        counts_t.plot(kind="bar", stacked=True, ax=ax, color=[COLORS[c] for c in ORDER])
        ax.set_xlabel("")
        ax.tick_params(axis="x", labelsize=8, rotation=45)
        st.pyplot(fig)
        plt.close(fig)
    else:
        st.caption("butuh komentar dari beberapa periode untuk tren")

    st.subheader("kata kunci per label")
    lab = st.selectbox("label", ORDER)
    text = " ".join(df.loc[df["label"] == lab, "tokens"].fillna(""))
    if text.strip():
        cmap = {"positive": "Greens", "neutral": "Greys", "negative": "Reds"}[lab]
        wc = WordCloud(width=900, height=360, background_color="white", collocations=False,
                       colormap=cmap, random_state=7, max_words=80) \
            .generate_from_frequencies(pd.Series(text.split()).value_counts().to_dict())
        fig, ax = plt.subplots(figsize=(9.5, 3.6), constrained_layout=True)
        ax.imshow(wc, interpolation="bilinear")
        ax.axis("off")
        st.pyplot(fig)
        plt.close(fig)
    else:
        st.caption("tidak ada token untuk label ini")

    st.subheader("evaluasi model (bert vs non-bert)")
    metrics = read_csv("data/metrics.csv")
    if metrics is not None:
        st.dataframe(metrics.fillna(""), width="stretch")
        st.caption("svm dilatih pada label indoBERT berconfidence tinggi lalu diuji di test split; "
                   "indoBERT diukur pada test split publik SmSA (ulasan berlabel manusia)")
    else:
        st.info("metrik muncul setelah sentiment.py berjalan dengan cukup komentar berlabel yakin")

    absa = read_csv("data/aspect_summary.csv")
    if absa is not None:
        st.subheader("sentimen per aspek (absa)")
        st.dataframe(absa, width="stretch")

    pipeline_figs = [f for f in ["sentiment_overview.png", "sentiment_trend.png", "wordclouds.png",
                                 "absa.png", "topics.png", "accounts.png", "media.png"]
                     if os.path.exists(f"figures/{f}")]
    if pipeline_figs:
        with st.expander("grafik lain dari pipeline", expanded=False):
            for f in pipeline_figs:
                st.image(f"figures/{f}")

    st.subheader("penjelajah data")
    flab = st.multiselect("filter label", ORDER, default=ORDER)
    view = df[df["label"].isin(flab)] if flab else df
    st.dataframe(view[["date", "source", "author", "label", "confidence", "text"]]
                 .reset_index(drop=True), width="stretch")


def collect_stats():
    df = read_csv("data/results.csv")
    if df is None:
        return None
    n = len(df)
    vc = df["label"].value_counts()
    stats = {
        "jumlah_komentar_organik": n,
        "persen": {lab: round(100 * vc.get(lab, 0) / n, 1) for lab in ORDER},
    }
    clean = read_csv("data/clean.csv")
    if clean is not None and "is_bot" in clean:
        stats["bot_difilter"] = int(clean["is_bot"].sum())
    absa = read_csv("data/aspect_summary.csv")
    if absa is not None and len(absa):
        top = absa.loc[absa["total"].idxmax()]
        stats["aspek_terbahas"] = {
            "aspek": top["aspek"], "total": int(top["total"]),
            "positive": int(top["positive"]), "neutral": int(top["neutral"]),
            "negative": int(top["negative"]),
        }
    arts = read_csv("data/news_articles.csv")
    if arts is not None:
        stats["artikel_suara_com"] = len(arts)
    metrics = read_csv("data/metrics.csv")
    if metrics is not None:
        stats["metrik_evaluasi"] = metrics.fillna("").to_dict("records")
    return stats


def llm_insights(stats, api_key, base_url, model):
    try:
        r = requests.post(f"{base_url.rstrip('/')}/chat/completions",
                          headers={"Authorization": f"Bearer {api_key}"},
                          json={"model": model, "temperature": 0.4, "messages": [
                              {"role": "system", "content":
                                  "kamu analis social media. balas hanya dengan 3 sampai 5 bullet "
                                  "insight singkat bahasa indonesia untuk slide kesimpulan, tanpa pembuka."},
                              {"role": "user", "content":
                                  "data hasil analisis sentimen mixue indonesia (json):\n"
                                  + json.dumps(stats, ensure_ascii=False, default=str)
                                  + "\n\ntulis insight & rekomendasi bisnis."},
                          ]}, timeout=90)
        r.raise_for_status()
        content = r.json()["choices"][0]["message"]["content"]
        lines = [ln.strip(" -•\t") for ln in content.splitlines() if ln.strip(" -•\t")]
        return lines or None
    except Exception:
        return None


def template_insights(stats):
    # fallback tanpa api key: narasi langsung dari angka pipeline, bukan teks kosong
    p = stats["persen"]
    dom = max(p, key=p.get)
    lines = [f"sentimen dominan {dom} ({p[dom]}% dari {stats['jumlah_komentar_organik']:,} "
             f"komentar organik; positif {p['positive']}%, negatif {p['negative']}%)"]
    if "aspek_terbahas" in stats:
        a = stats["aspek_terbahas"]
        lines.append(f"aspek paling banyak dibahas: {a['aspek']} ({a['total']} komentar) dengan "
                     f"komposisi {a['positive']} positif / {a['neutral']} netral / {a['negative']} negatif")
    if "artikel_suara_com" in stats:
        lines.append(f"media: {stats['artikel_suara_com']} artikel mixue di suara.com masuk periode data")
    for row in stats.get("metrik_evaluasi", []):
        if row.get("kelas") == "semua (smssa test)":
            lines.append(f"indoBERT mencapai akurasi {row['precision']} pada benchmark publik SmSA; "
                         "transformer lebih andal untuk slang daripada tf-idf + svm")
        if row.get("kelas") == "kappa kesesuaian":
            lines.append(f"kappa kesesuaian svm vs indoBERT {row['precision']} - model klasik cukup "
                         "sebagai pembanding murah tapi tidak menggantikan transformer")
    if p["negative"] > p["positive"]:
        lines.append("rekomendasi: prioritaskan perbaikan pada aspek dengan porsi negatif tertinggi "
                     "dan pantau komplain berulang di kolom komentar")
    else:
        lines.append("rekomendasi: pertahankan kualitas yang sudah disukai publik dan dorong konten "
                     "dari komentator paling aktif (kandidat kol)")
    return lines


def build_ppt(api_key, base_url, model):
    from pptx import Presentation
    from pptx.util import Inches, Pt

    stats = collect_stats()
    if stats is None:
        raise SystemExit("belum ada data/results.csv")

    insights = llm_insights(stats, api_key, base_url, model) if api_key else None

    prs = Presentation()

    def blank():
        return prs.slides.add_slide(prs.slide_layouts[6])

    def title(slide, text, size=30):
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.9))
        p = box.text_frame.paragraphs[0]
        p.text = text
        p.font.size = Pt(size)
        p.font.bold = True

    def bullets(slide, lines, left, top, width, height, size=15):
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        for i, line in enumerate(lines):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = "- " + str(line)
            para.font.size = Pt(size)
        return box

    def pic(slide, path, left, top, width):
        if os.path.exists(path):
            slide.shapes.add_picture(path, left, top, width=width)

    s = blank()
    title(s, "Analisis Sentimen Mixue Indonesia", 36)
    bullets(s, ["sumber: komentar instagram (akun resmi + hashtag) & portal berita suara.com",
                "pelabelan: indoBERT (pre-trained) lokal, pembanding tf-idf + svm",
                "preprocessing 2 tahap + filter bot/spam sesuai brief"], Inches(0.5), Inches(1.6),
            Inches(9), Inches(2.5), 18)

    s = blank()
    title(s, "metodologi")
    bullets(s, ["crawling instagram via instaloader (non-API): akun resmi mixueindonesia, "
                "hashtag #mixueindonesia & #mixue, komentar termasuk balasan",
                "crawling suara.com: artikel tag mixue + komentar pembaca dari endpoint portal",
                "stage 01: case folding, noise removal, slang, stopwords + stemming sastrawi",
                "stage 02: deteksi bot/spam berbasis aturan (akun toko, promosi, duplikat)",
                "sentimen indoBERT w11wo (lokal), absa per aspek, lda topic modeling",
                "evaluasi: svm (tf-idf) di test split + indoBERT di benchmark publik SmSA"],
            Inches(0.5), Inches(1.3), Inches(9), Inches(4.8), 15)

    p = stats["persen"]
    s = blank()
    title(s, "hasil sentimen")
    bullets(s, [f"komentar organik: {stats['jumlah_komentar_organik']:,}",
                f"positif {p['positive']}% | netral {p['neutral']}% | negatif {p['negative']}%"],
            Inches(0.5), Inches(1.2), Inches(4.3), Inches(1.6), 16)
    pic(s, "figures/sentiment_overview.png", Inches(0.5), Inches(2.7), Inches(9))

    s = blank()
    title(s, "tren & kata kunci")
    pic(s, "figures/sentiment_trend.png", Inches(0.5), Inches(1.25), Inches(9))
    pic(s, "figures/wordclouds.png", Inches(0.9), Inches(4.0), Inches(8.2))

    s = blank()
    title(s, "evaluasi: bert vs non-bert")
    if stats.get("metrik_evaluasi"):
        rows = stats["metrik_evaluasi"]
        table = s.shapes.add_table(len(rows) + 1, 5, Inches(0.5), Inches(1.3),
                                   Inches(9), Inches(0.4 * (len(rows) + 1))).table
        for j, col in enumerate(["model", "kelas", "precision", "recall", "f1"]):
            table.cell(0, j).text = col
        for i, row in enumerate(rows, start=1):
            for j, col in enumerate(["model", "kelas", "precision", "recall", "f1"]):
                table.cell(i, j).text = str(row.get(col, ""))
    else:
        bullets(s, ["belum ada metrik - jalankan sentiment.py dengan data cukup"],
                Inches(0.5), Inches(1.5), Inches(9), Inches(1))

    s = blank()
    title(s, "absa: sentimen per aspek")
    if "aspek_terbahas" in stats:
        a = stats["aspek_terbahas"]
        bullets(s, [f"aspek terbahas: {a['aspek']} ({a['total']} komentar)",
                    f"komposisi: {a['positive']} positif / {a['neutral']} netral / {a['negative']} negatif"],
                Inches(0.5), Inches(1.2), Inches(9), Inches(1.2), 16)
    pic(s, "figures/absa.png", Inches(0.5), Inches(2.5), Inches(9))

    s = blank()
    title(s, "topik & analisis akun/media")
    pic(s, "figures/topics.png", Inches(0.5), Inches(1.3), Inches(9))
    pic(s, "figures/accounts.png", Inches(0.5), Inches(4.3), Inches(4.4))
    pic(s, "figures/media.png", Inches(5.1), Inches(4.3), Inches(4.4))

    s = blank()
    title(s, "insight & kesimpulan")
    if insights is None:
        insights = template_insights(stats)
    bullets(s, insights, Inches(0.5), Inches(1.3), Inches(9), Inches(4.8), 16)

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


def page_ppt():
    st.title("generator ppt otomatis")
    st.caption("slide dibangun langsung dari angka & grafik pipeline. api key bersifat opsional: "
               "tanpa key, narasi insight dirangkai dari angka data; dengan key (openai-compatible), "
               "narasi dipoles model llm. key hanya dipakai di sesi ini, tidak disimpan.")
    df = read_csv("data/results.csv")
    if df is None:
        st.warning("belum ada data - jalankan pipeline dulu:")
        st.code(RUN_STEPS, language="bash")
        return

    api_key = st.text_input("api key llm (opsional)", type="password")
    base_url = st.text_input("base url openai-compatible", value="https://api.openai.com/v1")
    model = st.text_input("model", value="gpt-4o-mini")
    if st.button("buat ppt", type="primary"):
        try:
            with st.spinner("menyusun slide dari data..."):
                buf = build_ppt(api_key, base_url, model)
            st.download_button("download .pptx", buf, file_name="mixue_sentimen.pptx",
                               mime="application/vnd.openxmlformats-officedocument.presentationml.presentation")
            st.caption("untuk google slides: unggah file ke google drive lalu buka dengan slides, "
                       "atau di slides.google.com pilih File > Import slides.")
        except Exception as e:
            st.error(f"gagal membuat ppt: {type(e).__name__}: {e}")


page = st.sidebar.radio("menu", ["Dashboard", "Generator PPT"])
sidebar_status()
if page == "Dashboard":
    page_dashboard()
else:
    page_ppt()
